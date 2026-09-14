"""Exercise the browser's real contracts; Python is test-only, never application runtime."""
import io, json, os, re, warnings
from pathlib import Path
import requests
from openpyxl import Workbook, load_workbook
warnings.filterwarnings('ignore',message='Workbook contains no default style')

BASE=os.getenv('LOGICORE_TEST_URL','http://127.0.0.1:5000')
OUT=Path(__file__).parent/'generated'; OUT.mkdir(exist_ok=True)
s=requests.Session()
def check(r,status=200):
 assert r.status_code==status, (r.url,r.status_code,r.text[:1600])
 return r

def api(path,**kwargs):
 r=check(s.post(BASE+path,**kwargs) if kwargs else s.get(BASE+path))
 d=r.json();assert d.get('ok') is True,(path,d);return d

def book(sheets):
 w=Workbook();w.remove(w.active)
 for name,(headers,rows) in sheets.items():
  ws=w.create_sheet(name);ws.append(headers)
  for row in rows: ws.append(row)
 b=io.BytesIO();w.save(b);return b.getvalue()

def upload(path,filename,data,fields=None):
 return api(path,files={filename:(filename+'.xlsx',data)},data=fields or {})

def done(path):
 r=check(s.get(BASE+path));lines=[x[6:] for x in r.text.splitlines() if x.startswith('data: ')]
 d=json.loads(lines[-1]);assert d.get('type')=='done' and d.get('success') is True,d
 fn=d['filename'];r=check(s.get(BASE+'/download/'+fn));assert r.content[:2]==b'PK',r.text[:300];(OUT/fn).write_bytes(r.content)
 return d,load_workbook(io.BytesIO(r.content))

check(s.post(BASE+'/login',data={'username':'admin','password':'admin'},allow_redirects=False),302)
assert api('/healthz')['engine']=='Lucee'
for protected in ['/services/InvoiceService.cfc','/config/serial_rules.json','/template/FedEx_Shipment_Upload_Template.xlsx','/server.json']:check(s.get(BASE+protected),404)
portal=check(s.get(BASE+'/')).text
assert '"invoiceChildren"' in portal and '"isSuperadmin"' in portal
for p,active in [('/training-tracker/','portal-nav-training-tracker'),('/inventory-management/','portal-nav-inventory-management')]:
 page=check(s.get(BASE+p)).text
 assert 'class="app-shell"' in page and 'id="sidebar"' in page,page[:1600]
 assert f'id="{active}"' in page and f'id="{active}" href' in page,page[:1600]
 assert '/static/js/sidebar-nav.js' in page and 'localStorage.getItem("theme")' in page,page[:1600]
 assert 'â' not in page and 'Â' not in page,page[:1600]
check(s.get(BASE+'/admin/permissions'))

# Config round-trip tests the actual upload/download argument order and numeric data.
dims=book({'Dimensions':(['Model','Sq Footage'],[['TEST-75',20],['TEST-86',35]])})
for client in ['amc','philips']:
 upload('/upload_'+client+'_dimensions','file',dims)
 r=check(s.get(BASE+'/download_'+client+'_dimensions'))
 w=load_workbook(io.BytesIO(r.content));assert w.active['B2'].value==20
 assert api('/get_'+client+'_dimensions')['dimensions']['TEST-75']==20

# AMC: JSON manual correction must persist and trigger fresh analysis.
files={
 'receiving':('receiving.xlsx',book({'Receiving Export':(['Model','Serial Number','Received Date'],[['TEST-75','R1','08-03-2026']])})),
 'shipping':('shipping.xlsx',book({'Shipping Export':(['Model','Serial Number','Shipped Date'],[['TEST-75','S1','08-04-2026']])})),
 'inventory':('inventory.xlsx',book({'Inventory Export':(['Model','Serial Number'],[['MISSING','I1'],['MISSING','I1'],['TEST-75','I2']])}))}
d=api('/analyze_amc',files=files,data={'period_start':'2026-08-01','period_end':'2026-08-31','invoice_title':'AMC parity title','output_filename':'AMC_PARITY'})
assert d['missing_dimension_models']==['MISSING'],d
api('/confirm_amc',json={'dimensions':{'MISSING':1000}})
d,w=done('/stream_amc');assert d['filename']=='AMC_PARITY.xlsx';assert d['additional_sqft']==500,d;assert d['subtotal']==3686,d

# Philips: missing Size column is valid; manual dimensions and metadata survive confirm.
p=book({'Inventory':(['Model','Serial','Type'],[['MISSING-PH','I1','Demo'],['TEST-75','I2','Service']]),'Shipping':(['Model','Stock Level (Primary)'],[['TEST-75','Service']]),'Recieved':(['Model'],[['TEST-75']]),'Repairs':(['Model','Status'],[['75BDL','Repaired']])})
d=upload('/analyze_philips','report',p,{'output_filename':'PHILIPS_PARITY','invoice_title':'Philips parity title','parts_sqft':'100'})
assert 'MISSING-PH' in d['missing_dimension_models']
api('/confirm_philips',json={'dimensions':{'MISSING-PH':600}})
d,w=done('/stream_philips');assert d['filename']=='PHILIPS_PARITY.xlsx';assert d['demo_total_sqft']==600,d

# TCL must consume the JSON payload sent by the unchanged frontend.
p=book({'Inventory Export':(['Model','Serial Number','Grade','Rack','Bin','Received Date'],[['TV','U1','A','MAIN','X','08-01-2026'],['PART','P1','A','PARTS','X','08-02-2026']])})
d=upload('/analyze_tcl','inventory',p,{'date_from':'2026-08-01','date_to':'2026-08-31','output_filename':'TCL_PARITY'})
api('/confirm_tcl',json={'unit_breakdowns':{d['unit_groups'][0]['key']:'1'},'box_breakdowns':{d['part_groups'][0]['key']:'1'}})
d,w=done('/stream_tcl');assert d['filename']=='TCL_PARITY.xlsx';assert d['subtotal']==78.85,d

print('PASS: JSON casing, portal navigation, dimensions round-trip, AMC/Philips correction recomputation, TCL JSON confirmation and XLSX downloads')

# Worksheet contracts are part of parity, including the intentionally misspelled source name.
w=load_workbook(OUT/'AMC_PARITY.xlsx');assert w.sheetnames==['Breakdown','Received','Shipped','Excluded Items'];assert w['Breakdown']['A5'].value=='AMC parity title';assert w['Breakdown']['D8'].value=='=COUNTA(Received!C2:C1048576)';assert w['Received']['C2'].value.day==3
w=load_workbook(OUT/'PHILIPS_PARITY.xlsx');assert w['Breakdown']['D16'].value=='=COUNTA(Shipping!G2:G100000)';assert w['Shipping']['G2'].value=='TEST-75';assert w['Repairs']['G2'].value=='Yes'
w=load_workbook(OUT/'TCL_PARITY.xlsx');assert w.sheetnames==['Invoice','Line Items'];assert w['Line Items']['G2'].value==75

# FedEx preserves template columns (including duplicate MainContact), real Excel dates/time,
# truncation instead of rounding, blank-recipient defaults, and tracking identifiers.
fheaders=['Express or Ground Tracking ID','Net Charge Amount','Payor','Recipient Company','Recipient Name','Recipient Address Line 1','Recipient Address Line 2','Recipient City','Recipient State','Recipient Zip Code','Original Customer Reference','Original Ref#3/PO Number']
p=book({'Raw':(fheaders,[['001234567890123456',8.5,'','','','','','','','','PO1',''],['001234567890123456',8.5,'','','','','','','','','',''],['T2',1,'','Example','Matt Shaw','','','Melbourne','fl','32940-1234','','PO2'],['T3',0,'','','','','','','','','','']])})
d=upload('/analyze_fedex_shipment','raw',p,{'period_label':'082026','call_date':'2026-08-31','output_filename':'FEDEX_PARITY'})
api('/build_fedex_shipment',json={});d,w=done('/stream_fedex_shipment');assert d['row_count']==2,d;assert d['total_price']==11.17,d;assert d['skipped_count']==2,d
headers=[c.value for c in w.active[1]];assert len(headers)>23;assert headers.count('MainContact')==2
assert w.active.cell(2,headers.index('Summary')+1).value.endswith('001234567890123456')
assert w.active.cell(2,headers.index('CallRcvd')+1).value.day==31
assert w.active.cell(2,headers.index('CallRcvdTime')+1).value.hour==17

# Workshop: previously triaged repairs get the reduced rate; previously repaired
# serials require a later shipment; source/master companion files retain other sheets.
raw=book({'Repair Data':(['Date Integer','Actual Model','Actual Serial','Derive Size','Result','Category'],[
 ['08-05-2026','AP9-A75-NA-R','9A75XK001','75','Mainboard replaced','Refurbished'],
 ['08-05-2026','AP9-A75-NA-R','9A75XK002','75','Mainboard replaced','Refurbished'],
 ['08-06-2026','AP9-A86-NA-R','9A86XK003','86','Pending LCD','Pending Parts'],
 ['07-01-2026','AP9-A75-NA-R','9A75XK004','75','Mainboard replaced','Refurbished']]),'Notes':(['Keep'],[['untouched']])})
master=book({'Repair Log':(['Month','Model','Serial','Type'],[['2026-07-01','AP9-A75-NA-R','9A75XK002','Basic']]),'Triage Log':(['Date','Model','Serial','Type'],[['2026-07-01','AP9-A75-NA-R','9A75XK001','Triage - Basic']]),'Notes':(['Keep'],[['history']])})
ship=b'Shipped Date,Serial Number\n07-01-2026,9A75XK002\n'
d=api('/sanitize',files={'raw_file':('raw.xlsx',raw),'prev_invoiced':('master.xlsx',master),'shipping':('ship.csv',ship)},data={'date_from':'2026-08-01','date_to':'2026-08-31','invoice_date':'2026-09-01','completed_date':'2026-08-31','customer':'Promethean','call_id':'C-PARITY'})
assert d['total_records']==3,d;assert d['issue_count']==0,d
api('/generate',data={'output_filename':'WORKSHOP_PARITY','corrections':'{}'})
d,w=done('/stream');assert d['subtotal']==245,d;assert d['excluded_count']==1,d
assert w['Breakdown']['E1'].value=='C-PARITY';assert w['Depot Repair']['C2'].value=='BasicSmall - Previously Triaged';assert w['Depot Repair']['E2'].value==64
assert w['Triage Units']['D2'].value=='HeavyLarge-Triage';assert 'TblPartTesting' in w['Part Testing & Programming'].tables
for field in ['corrected_filename','master_filename']:
 data=check(s.get(BASE+'/download/'+d[field])).content;cw=load_workbook(io.BytesIO(data));assert 'Notes' in cw.sheetnames
 if field=='master_filename':assert cw['Repair Log'].max_row==3;assert cw['Triage Log'].max_row==3
 else:assert 'Sanitized Data' in cw.sheetnames

# NonConforming CRUD and export use JSON and retain punctuation in cell values.
item=api('/nonconforming/api/items',json={'model':'Parity, Model','serial':'001234','carrier':'FedEx','status':'Pending'})['item'];iid=item['id']
r=check(s.patch(BASE+f'/nonconforming/api/items/{iid}',json={'status':'Resolved','addtl_info':'First, second'}));assert r.json()['item']['status']=='Resolved',r.text
r=check(s.get(BASE+'/nonconforming/api/export',params={'q':'Parity, Model'}));w=load_workbook(io.BytesIO(r.content));assert any(row[3]=='Parity, Model' and row[4]=='001234' for row in w.active.iter_rows(min_row=2,values_only=True))
label=api(f'/nonconforming/api/items/{iid}/label');assert '^XA' in label['zpl'] and item['number'] in label['zpl']
check(s.delete(BASE+f'/nonconforming/api/items/{iid}'));check(s.get(BASE+f'/nonconforming/api/items/{iid}'),404)
print('PASS: source workbook layouts/formulas, FedEx template/date/time, Workshop dedup/rates/companions, NonConforming CRUD/export/labels')

# Training: create content and attendance, then exercise signed scans and reports.
def postform(path,data):return check(s.post(BASE+path,data=data))
r=postform('/training-tracker/week/new',{'title':'Parity Week','start_date':'2026-08-03','notes':'Test notes'});wid=re.search(r'/week/(\d+)',r.url).group(1)
postform('/training-tracker/people/add',{'name':'Parity Person','email':'parity@example.test'})
postform(f'/training-tracker/week/{wid}/topic/new',{'title':'Parity Topic','lesson_plan':'Test lesson','key_points':'One\nTwo'})
r=postform(f'/training-tracker/week/{wid}/session/new',{'trainer_name':'Parity Trainer','session_date':'2026-08-03','location':'Workshop'})
sid=re.findall(r'href="/training-tracker/session/(\d+)"',r.text)[-1]
r=check(s.get(BASE+f'/training-tracker/session/{sid}'));pid=re.findall(r'/attendance/(\d+)"',r.text)[-1]
postform(f'/training-tracker/session/{sid}/attendance/{pid}',{'attended':'on','signature':'Parity Person'})
r=check(s.get(BASE+f'/training-tracker/session/{sid}/signoff/download'));assert r.content.startswith(b'%PDF'),r.text[:500]
pdf=r.content
postform(f'/training-tracker/session/{sid}/attendance/toggle',{'person_id':pid,'attended':'1'})
r=check(s.post(BASE+f'/training-tracker/session/{sid}/signoff/upload',files={'signoff_file':('signed.pdf',pdf,'application/pdf')}))
r=check(s.get(BASE+f'/training-tracker/session/{sid}/signoff/view'));assert r.content==pdf
r=check(s.get(BASE+'/training-tracker/reports/attendance-matrix',params={'start_week':wid,'end_week':wid}));w=load_workbook(io.BytesIO(r.content));assert any('X' in row for row in w.active.iter_rows(values_only=True))
r=check(s.get(BASE+'/training-tracker/reports/signoff-zip',params={'start_week':wid,'end_week':wid}));assert r.content.startswith(b'PK')
postform(f'/training-tracker/week/{wid}/delete',{});postform(f'/training-tracker/people/{pid}/delete',{})

# Inventory imports exercise rows with no MSO and full FedEx event details.
import uuid
serial='PARITY'+uuid.uuid4().hex[:8]
p=book({'Receiving Export':(['Received Date','Model','Serial Number','RMA In Tracking'],[['08-03-2026','AP9-A75-NA-R',serial,'']]),
 'Shipping Export':(['Ticket Number','Shipped Date','Model','Serial Number','Tracking Number'],[['M99999999','08-08-2026','AP9-A75-NA-R',serial,'00123456789']]),
 'FedEx Master':(['MSO','Serial Number','Outbound Tracking','Request Date','Fedex Ship Date','Outbound Delivery Date','Return Tracking','Return Tracking Ship Date','Return Tracking Delivery Date'],[['M99999999',serial,'OUT','08-01-2026','08-02-2026','08-03-2026','RETURN','08-04-2026','08-05-2026']])})
r=check(s.post(BASE+'/inventory-management/import',files={'files':('parity.xlsx',p)}));assert '8 events' in r.text,r.text[-1600:]
r=check(s.post(BASE+'/inventory-management/import',files={'files':('parity.xlsx',p)}));assert 'Duplicate file' in r.text
r=check(s.get(BASE+'/inventory-management/serial/'+serial));assert 'fedex_return_delivered' in r.text
r=check(s.get(BASE+'/inventory-management/shipping/all/export.xlsx',params={'preset':'custom','start':'2026-08-01','end':'2026-08-31'}));w=load_workbook(io.BytesIO(r.content));assert any(serial in row for row in w.active.iter_rows(values_only=True))
print('PASS: Training content/attendance/PDF/upload/ZIP/XLSX, Inventory imports/dedup/lifecycle/export')

# The source portal exposes a second Workshop path for pre-sanitized files.
legacy=book({'Repair Data':(['Date Integer','Actual Model','Actual Serial','Derive Size','Type','Type2','Result','Category'],[['08-10-2026','AP9-A75-NA-R','LEGACY-PARITY','75','Depot Repair Tab','Basic','Repaired','Refurbished']])})
r=api('/generate',files={'repair':('legacy.xlsx',legacy),'prev_invoiced':('master.xlsx',master),'shipping':('ship.csv',b'Shipped Date,Serial Number\n')},data={'mode':'legacy','date_from':'2026-08-01','date_to':'2026-08-31','invoice_date':'2026-09-01','completed_date':'2026-08-31','customer':'Promethean','call_id':'C-LEGACY','output_filename':'LEGACY_PARITY'})
d,w=done('/stream');assert d['filename']=='LEGACY_PARITY.xlsx';assert d['subtotal']==110,d;assert w['Breakdown']['E1'].value=='C-LEGACY';assert 'corrected_filename' not in d

# Limited users only see and reach explicitly granted sections.
username='limited_'+uuid.uuid4().hex[:8]
postform('/admin/permissions/users/new',{'username':username,'password':'parity-pass','initials':'LP'})
admin_page=check(s.get(BASE+'/admin/permissions')).text
card=re.search(r'<div class="card"><h2>'+username+r'.*?action="/admin/permissions/users/(\d+)/set"',admin_page,re.S);assert card,username
uid=card.group(1);postform(f'/admin/permissions/users/{uid}/set',{'ig_amc':'on','sms_nonconforming':'on','training_role':'viewer'})
limited=requests.Session();check(limited.post(BASE+'/login',data={'username':username,'password':'parity-pass'},allow_redirects=False),302)
portal=check(limited.get(BASE+'/')).text;assert '"invoiceChildren":["amc"]' in portal;assert '"trainingRole":"viewer"' in portal
check(limited.post(BASE+'/analyze_amc'),400);check(limited.get(BASE+'/get_philips_dimensions'),403);assert check(limited.get(BASE+'/inventory-management/')).url==BASE+'/'
check(limited.get(BASE+'/training-tracker/'));check(limited.get(BASE+'/training-tracker/admin'),403)
postform(f'/admin/permissions/users/{uid}/delete',{})
print('PASS: legacy Workshop workflow and restricted-user section/role permissions')
